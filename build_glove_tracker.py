"""Builds glove_tracker.xlsx — a red-themed glove stock & issue tracker.

The workbook is designed to be driven by the macros in GloveTracker.bas
(Issue Gloves / Undo Last Issue / Add Colleague buttons on the Dashboard).

Run:  python build_glove_tracker.py
"""
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "glove_tracker.xlsx"
FONT = "Arial"

# Red theme palette
DARK_RED = "7F0000"
RED = "C00000"
MID_RED = "E06666"
LIGHT_RED = "F4CCCC"
PALE = "FFF5F5"
WHITE = "FFFFFF"
GREY = "595959"

SIZES = ["Small", "Medium", "Large"]
N_COLLEAGUES = 200   # rows available on Colleagues sheet  (5..204)
N_LOG = 1000         # rows available on Issue Log         (5..1004)
DATE_FMT = "dd-mmm-yyyy"
UNLOCKED = Protection(locked=False)

thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def title(ws, text, subtitle, width_cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols)
    c = ws.cell(1, 1, text)
    c.font = Font(name=FONT, size=18, bold=True, color=WHITE)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 36
    for col in range(1, width_cols + 1):
        ws.cell(1, col).fill = fill(DARK_RED)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
    s = ws.cell(2, 1, subtitle)
    s.font = Font(name=FONT, size=10, italic=True, color=GREY)
    s.alignment = Alignment(indent=1, wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30
    ws.sheet_view.showGridLines = False


def header_row(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row, start_col + i, h)
        c.font = Font(name=FONT, bold=True, color=WHITE)
        c.fill = fill(RED)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 30


def body_cell(c, input_cell=False, fmt=None, center=False):
    c.font = Font(name=FONT, color="000000")
    c.border = BORDER
    if input_cell:
        c.fill = fill(PALE)
        c.protection = UNLOCKED
    if fmt:
        c.number_format = fmt
    if center:
        c.alignment = Alignment(horizontal="center")


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def protect(ws):
    # No password: Review > Unprotect Sheet works if a manual fix is ever needed.
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.autoFilter = False


wb = Workbook()
dash = wb.active
dash.title = "Dashboard"
col_ws = wb.create_sheet("Colleagues")
log = wb.create_sheet("Issue Log")
stock = wb.create_sheet("Stock")
setup = wb.create_sheet("Macro Setup")
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = RED

size_dv_formula = '"' + ",".join(SIZES) + '"'

# ---------------------------------------------------------------- Colleagues
C_FIRST, C_LAST = 5, 5 + N_COLLEAGUES - 1
title(col_ws, "COLLEAGUES",
      "Type names in column B and pick each person's glove size in column C (pink cells), "
      "or use the Add Colleague button on the Dashboard. Everything else fills in automatically. "
      "3rd pair or more = FLAG.", 10)
header_row(col_ws, 4, ["#", "Name", "Glove Size", "Total Pairs Given",
                       "1st Pair Date", "2nd Pair Date", "3rd Pair Date",
                       "Last Issued", "Status", "Flag #"])
L = "'Issue Log'!"
LD = f"{L}$A$5:$A${4 + N_LOG}"
LN = f"{L}$B$5:$B${4 + N_LOG}"
LS = f"{L}$C$5:$C${4 + N_LOG}"
LQ = f"{L}$D$5:$D${4 + N_LOG}"
LR = f"{L}$E$5:$E${4 + N_LOG}"
CN = f"Colleagues!$B${C_FIRST}:$B${C_LAST}"
CS = f"Colleagues!$C${C_FIRST}:$C${C_LAST}"
for r in range(C_FIRST, C_LAST + 1):
    col_ws.cell(r, 1, r - C_FIRST + 1)
    body_cell(col_ws.cell(r, 1), center=True)
    col_ws.cell(r, 1).font = Font(name=FONT, color=GREY)
    body_cell(col_ws.cell(r, 2), input_cell=True)
    body_cell(col_ws.cell(r, 3), input_cell=True, center=True)
    col_ws.cell(r, 4, f'=IF(B{r}="","",SUMIFS({LQ},{LN},B{r}))')
    body_cell(col_ws.cell(r, 4), center=True)
    for n, col in ((1, 5), (2, 6), (3, 7)):
        col_ws.cell(r, col,
                    f'=IF(OR(B{r}="",N(D{r})<{n}),"",_xlfn.MINIFS({LD},{LN},B{r},{LR},">="&{n}))')
        body_cell(col_ws.cell(r, col), fmt=DATE_FMT, center=True)
    col_ws.cell(r, 8, f'=IF(N(D{r})=0,"",_xlfn.MAXIFS({LD},{LN},B{r}))')
    body_cell(col_ws.cell(r, 8), fmt=DATE_FMT, center=True)
    col_ws.cell(r, 9, f'=IF(B{r}="","",IF(D{r}>=3,"FLAG",IF(D{r}=0,"Not issued","OK")))')
    body_cell(col_ws.cell(r, 9), center=True)
    col_ws.cell(r, 10, f'=IF(I{r}="FLAG",COUNTIF($I${C_FIRST}:I{r},"FLAG"),"")')
    body_cell(col_ws.cell(r, 10), center=True)
    col_ws.cell(r, 10).font = Font(name=FONT, color=GREY, size=8)

widths(col_ws, {"A": 5, "B": 28, "C": 12, "D": 12, "E": 14, "F": 14,
                "G": 14, "H": 14, "I": 12, "J": 7})
col_ws.freeze_panes = "C5"
dv = DataValidation(type="list", formula1=size_dv_formula, allow_blank=True)
dv.add(f"C{C_FIRST}:C{C_LAST}")
col_ws.add_data_validation(dv)
col_ws.conditional_formatting.add(
    f"A{C_FIRST}:J{C_LAST}", FormulaRule(formula=[f'$I{C_FIRST}="FLAG"'], fill=fill(LIGHT_RED),
                                         font=Font(name=FONT, color=DARK_RED, bold=True)))
col_ws.conditional_formatting.add(
    f"I{C_FIRST}:I{C_LAST}",
    CellIsRule(operator="equal", formula=['"FLAG"'], fill=fill(RED),
               font=Font(name=FONT, color=WHITE, bold=True)))
col_ws["J4"].comment = Comment("Helper number used by the Dashboard's flagged list.", "Glove Tracker")
protect(col_ws)

# ---------------------------------------------------------------- Issue Log
G_FIRST, G_LAST = 5, 4 + N_LOG
title(log, "ISSUE LOG",
      "Filled in by the Issue Gloves button on the Dashboard (which also takes the gloves off stock). "
      "Sheet is locked so entries always go through the button.", 6)
header_row(log, 4, ["Date Given", "Name", "Size", "Pairs Given",
                    "Pair # (running total)", "Flag"])
for r in range(G_FIRST, G_LAST + 1):
    body_cell(log.cell(r, 1), fmt=DATE_FMT, center=True)
    body_cell(log.cell(r, 2))
    body_cell(log.cell(r, 3), center=True)
    body_cell(log.cell(r, 4), center=True)
    log.cell(r, 5, f'=IF(B{r}="","",SUMIFS($D${G_FIRST}:D{r},$B${G_FIRST}:B{r},B{r}))')
    body_cell(log.cell(r, 5), center=True)
    log.cell(r, 6, f'=IF(B{r}="","",IF(E{r}>=3,"FLAG",""))')
    body_cell(log.cell(r, 6), center=True)
widths(log, {"A": 14, "B": 28, "C": 11, "D": 11, "E": 14, "F": 10})
log.freeze_panes = "A5"
log.conditional_formatting.add(
    f"A{G_FIRST}:F{G_LAST}", FormulaRule(formula=[f'$F{G_FIRST}="FLAG"'], fill=fill(LIGHT_RED),
                                         font=Font(name=FONT, color=DARK_RED, bold=True)))
log.conditional_formatting.add(
    f"F{G_FIRST}:F{G_LAST}",
    CellIsRule(operator="equal", formula=['"FLAG"'], fill=fill(RED),
               font=Font(name=FONT, color=WHITE, bold=True)))
protect(log)

# ---------------------------------------------------------------- Stock
title(stock, "STOCK",
      "Type how many pairs you have of each size in On Hand (pink cells). When new gloves arrive, "
      "change the number. The Issue Gloves button takes pairs off automatically.", 5)
header_row(stock, 4, ["Size", "On Hand", "Issued (all time)", "Status"])
for i, size in enumerate(SIZES):
    r = 5 + i
    stock.cell(r, 1, size)
    stock.cell(r, 2, 0)
    stock.cell(r, 3, f"=SUMIFS({LQ},{LS},A{r})")
    stock.cell(r, 4, f'=IF(B{r}<=0,"OUT OF STOCK",IF(B{r}<$B$10,"LOW","OK"))')
    for c in range(1, 5):
        body_cell(stock.cell(r, c), input_cell=(c == 2), center=c > 1)
    stock.cell(r, 1).font = Font(name=FONT, bold=True)
stock.cell(8, 1, "Total")
stock.cell(8, 2, "=SUM(B5:B7)")
stock.cell(8, 3, "=SUM(C5:C7)")
for c in range(1, 5):
    body_cell(stock.cell(8, c), center=c > 1)
    stock.cell(8, c).font = Font(name=FONT, bold=True)
    stock.cell(8, c).fill = fill(LIGHT_RED)

stock.cell(10, 1, "Low-stock alert below")
stock.cell(10, 1).font = Font(name=FONT, bold=True)
stock.cell(10, 2, 10)
body_cell(stock.cell(10, 2), input_cell=True, center=True)
stock.cell(10, 3, "pairs per size (agreed default, change as needed)")
stock.cell(10, 3).font = Font(name=FONT, italic=True, color=GREY)
dv_q = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0",
                      showErrorMessage=True, error="Enter a whole number of pairs (0 or more).")
dv_q.add("B5:B7")
dv_q.add("B10")
stock.add_data_validation(dv_q)
widths(stock, {"A": 22, "B": 14, "C": 18, "D": 16, "E": 40})
for cond, color in (('"OUT OF STOCK"', DARK_RED), ('"LOW"', MID_RED)):
    stock.conditional_formatting.add(
        "D5:D7", CellIsRule(operator="equal", formula=[cond], fill=fill(color),
                            font=Font(name=FONT, color=WHITE, bold=True)))
protect(stock)

# ---------------------------------------------------------------- Dashboard
title(dash, "GLOVE TRACKER — DASHBOARD",
      "Issue gloves and add colleagues with the panels below. Anyone on their 3rd pair or more is FLAGGED.", 12)
widths(dash, {c: 12 for c in "ABCDEFGHIJKL"})
dash.column_dimensions["A"].width = 3
dash.column_dimensions["F"].width = 15


def tile(row, col, label, formula, big_color=RED):
    dash.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    dash.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=col + 1)
    lab = dash.cell(row, col, label)
    lab.font = Font(name=FONT, bold=True, color=WHITE, size=10)
    lab.alignment = Alignment(horizontal="center", vertical="center")
    val = dash.cell(row + 1, col, formula)
    val.font = Font(name=FONT, bold=True, size=24, color=big_color)
    val.alignment = Alignment(horizontal="center", vertical="center")
    for r in range(row, row + 3):
        for c in range(col, col + 2):
            dash.cell(r, c).fill = fill(RED if r == row else PALE)
            dash.cell(r, c).border = BORDER


dash.row_dimensions[4].height = 22
tile(4, 2, "SMALL ON HAND", "=Stock!B5")
tile(4, 4, "MEDIUM ON HAND", "=Stock!B6")
tile(4, 6, "LARGE ON HAND", "=Stock!B7")
tile(4, 8, "TOTAL PAIRS ISSUED", "=Stock!C8", big_color=DARK_RED)
tile(4, 10, "FLAGGED PEOPLE", f'=COUNTIF(Colleagues!I{C_FIRST}:I{C_LAST},"FLAG")', big_color=DARK_RED)
FLAG_COUNT = "$J$5"

tile(8, 2, "COLLEAGUES LISTED", f'=COUNTA(Colleagues!B{C_FIRST}:B{C_LAST})', big_color=GREY)
tile(8, 4, "GIVEN 1 PAIR", f'=COUNTIF(Colleagues!D{C_FIRST}:D{C_LAST},1)', big_color=GREY)
tile(8, 6, "GIVEN 2 PAIRS", f'=COUNTIF(Colleagues!D{C_FIRST}:D{C_LAST},2)', big_color=GREY)
tile(8, 8, "NOT YET ISSUED", f'=COUNTIF(Colleagues!I{C_FIRST}:I{C_LAST},"Not issued")', big_color=GREY)
tile(8, 10, "FLAGGED HAND-OUTS", f'=COUNTIF({L}F5:F{4 + N_LOG},"FLAG")', big_color=DARK_RED)
for r in (5, 6, 9, 10):
    dash.row_dimensions[r].height = 22


def panel_header(row, c1, c2, text):
    dash.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    h = dash.cell(row, c1, text)
    h.font = Font(name=FONT, bold=True, color=WHITE, size=12)
    h.alignment = Alignment(indent=1, vertical="center")
    for c in range(c1, c2 + 1):
        dash.cell(row, c).fill = fill(DARK_RED)
    dash.row_dimensions[row].height = 24


def label(cell, text):
    dash[cell] = text
    dash[cell].font = Font(name=FONT, bold=True, color=DARK_RED)
    dash[cell].alignment = Alignment(horizontal="right", indent=1)


def entry(rng):
    first = rng.split(":")[0]
    if ":" in rng:
        dash.merge_cells(rng)
    cells = dash[rng] if ":" in rng else ((dash[rng],),)
    for row in cells:
        for c in row:
            c.fill = fill(WHITE)
            c.border = Border(left=Side(style="thin", color=RED), right=Side(style="thin", color=RED),
                              top=Side(style="thin", color=RED), bottom=Side(style="thin", color=RED))
            c.protection = UNLOCKED
    dash[first].font = Font(name=FONT, bold=True, size=11)
    dash[first].alignment = Alignment(horizontal="center", vertical="center")


# Issue Gloves panel  (cells read by the IssueGloves macro: C13 name, C14 size, C15 pairs)
panel_header(12, 2, 6, "ISSUE GLOVES")
label("B13", "Name")
entry("C13:D13")
label("B14", "Size")
entry("C14:D14")
label("B15", "Pairs")
entry("C15")
dash["C15"] = 1
dash["E14"] = (f'=IF($C$13="","",IFERROR(IF(INDEX({CS},MATCH($C$13,{CN},0))="","",'
               f'"Usual: "&INDEX({CS},MATCH($C$13,{CN},0))),""))')
dash["E14"].font = Font(name=FONT, italic=True, color=GREY, size=9)
dash["E13"] = "← pick from list"
dash["E13"].font = Font(name=FONT, italic=True, color=GREY, size=9)
dash["D15"] = "(blank size = usual size)"
dash["D15"].font = Font(name=FONT, italic=True, color=GREY, size=9)
dash.merge_cells("B16:F16")
had = f"SUMIFS({LQ},{LN},$C$13)"
dash["B16"] = (f'=IF($C$13="","Pick a name, then press ISSUE GLOVES.",'
               f'IF(ISNA(MATCH($C$13,{CN},0)),"Name not on Colleagues list",'
               f'"Has had "&{had}&" pair(s), this will be pair #"&({had}+IF(N($C$15)<1,1,N($C$15)))'
               f'&IF({had}+IF(N($C$15)<1,1,N($C$15))>=3,"  >> WILL BE FLAGGED","")))')
dash["B16"].font = Font(name=FONT, italic=True, color=GREY)
dash["B16"].alignment = Alignment(horizontal="center", vertical="center")
dash.row_dimensions[16].height = 22
dash.conditional_formatting.add(
    "B16:F16", FormulaRule(formula=['ISNUMBER(SEARCH("FLAGGED",$B$16))'], fill=fill(RED),
                           font=Font(name=FONT, color=WHITE, bold=True)))
for r in (17, 18):
    dash.row_dimensions[r].height = 18
# Buttons ISSUE GLOVES (B17:D18) and UNDO LAST ISSUE (E17:F18) are drawn by SetupButtons.

# Add Colleague panel  (cells read by the AddColleague macro: I13 name, I14 size)
panel_header(12, 8, 11, "ADD COLLEAGUE")
label("H13", "Name")
entry("I13:K13")
label("H14", "Size")
entry("I14")
# Button ADD COLLEAGUE (H16:K17) is drawn by SetupButtons.

dv_dname = DataValidation(type="list", formula1=f"={CN}", allow_blank=True, showErrorMessage=True,
                          errorTitle="Unknown name", error="Add this person with ADD COLLEAGUE first.")
dv_dname.add("C13")
dv_dsize = DataValidation(type="list", formula1=size_dv_formula, allow_blank=True)
dv_dsize.add("C14")
dv_dsize.add("I14")
dv_dqty = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                         showErrorMessage=True, error="Enter a whole number of pairs (1 or more).")
dv_dqty.add("C15")
for d in (dv_dname, dv_dsize, dv_dqty):
    dash.add_data_validation(d)

# Stock status table
ST = 20
header_row(dash, ST, ["Size", "On Hand", "Issued (all time)", "Status"], start_col=2)
for i in range(3):
    r = ST + 1 + i
    for j, col in enumerate("ABCD"):
        c = dash.cell(r, 2 + j, f"=Stock!{col}{5 + i}")
        body_cell(c, center=j > 0)
dash.cell(ST + 4, 2, "Total")
dash.cell(ST + 4, 3, "=Stock!B8")
dash.cell(ST + 4, 4, "=Stock!C8")
for c in range(2, 6):
    body_cell(dash.cell(ST + 4, c), center=c > 2)
    dash.cell(ST + 4, c).font = Font(name=FONT, bold=True)
    dash.cell(ST + 4, c).fill = fill(LIGHT_RED)
for cond, color in (('"OUT OF STOCK"', DARK_RED), ('"LOW"', MID_RED)):
    dash.conditional_formatting.add(
        f"E{ST + 1}:E{ST + 3}", CellIsRule(operator="equal", formula=[cond], fill=fill(color),
                                           font=Font(name=FONT, color=WHITE, bold=True)))
dash.cell(ST + 5, 2, '="Low-stock alert below "&Stock!B10&" pairs per size. Update On Hand on the Stock sheet."')
dash.cell(ST + 5, 2).font = Font(name=FONT, italic=True, color=GREY, size=9)

chart = BarChart()
chart.type = "col"
chart.title = "Stock by Size"
chart.y_axis.title = "Pairs"
chart.height, chart.width = 7, 11.5
data = Reference(stock, min_col=2, max_col=3, min_row=4, max_row=7)
cats = Reference(stock, min_col=1, min_row=5, max_row=7)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.series[0].graphicalProperties.solidFill = DARK_RED
chart.series[1].graphicalProperties.solidFill = MID_RED
chart.legend.position = "b"
dash.add_chart(chart, f"H{ST}")

# How-to
HT = ST + 7
dash.cell(HT, 2, "HOW TO USE").font = Font(name=FONT, bold=True, color=DARK_RED)
steps = [
    "1. Stock sheet: type how many pairs you have of each size. Change it when new gloves arrive.",
    "2. Add people with ADD COLLEAGUE (or type them on the Colleagues sheet).",
    "3. To give gloves: pick the name, size (blank = usual) and pairs, then press ISSUE GLOVES.",
    "4. Counts never reset. 1st & 2nd pair = OK; 3rd pair or more = FLAG (red).",
    "5. Made a mistake? UNDO LAST ISSUE removes the last hand-out and puts the stock back.",
    "   First time only: see the 'Macro Setup' sheet to switch the buttons on.",
]
for i, s in enumerate(steps):
    dash.cell(HT + 1 + i, 2, s).font = Font(name=FONT, size=9, color=GREY)

# Flagged list
FL = HT + 10
dash.merge_cells(start_row=FL - 1, start_column=2, end_row=FL - 1, end_column=10)
h = dash.cell(FL - 1, 2, "FLAGGED: 3rd PAIR OR MORE")
h.font = Font(name=FONT, bold=True, color=WHITE, size=12)
for c in range(2, 11):
    dash.cell(FL - 1, c).fill = fill(DARK_RED)
dash.row_dimensions[FL].height = 30
dash.merge_cells(start_row=FL, start_column=3, end_row=FL, end_column=4)
flag_headers = ["#", "Name", "Size", "Total Pairs", "1st Pair", "2nd Pair", "3rd Pair", "Last Issued"]
for col, text in zip("BCEFGHIJ", flag_headers):
    dash[f"{col}{FL}"] = text
    dash[f"{col}{FL}"].font = Font(name=FONT, bold=True, color=WHITE)
    dash[f"{col}{FL}"].fill = fill(RED)
    dash[f"{col}{FL}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
N_FLAG_ROWS = 25
src_cols = {"C": "B", "E": "C", "F": "D", "G": "E", "H": "F", "I": "G", "J": "H"}
for k in range(1, N_FLAG_ROWS + 1):
    r = FL + k
    dash.cell(r, 2, f'=IF({k}>{FLAG_COUNT},"",{k})')
    body_cell(dash.cell(r, 2), center=True)
    dash.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    for dcol, scol in src_cols.items():
        dash[f"{dcol}{r}"] = (f'=IF({k}>{FLAG_COUNT},"",INDEX(Colleagues!${scol}${C_FIRST}:${scol}${C_LAST},'
                              f'MATCH({k},Colleagues!$J${C_FIRST}:$J${C_LAST},0)))')
        body_cell(dash[f"{dcol}{r}"], fmt=DATE_FMT if dcol in "GHIJ" else None, center=dcol != "C")
    if k % 2 == 0:
        for c in "BCEFGHIJ":
            dash[f"{c}{r}"].fill = fill(PALE)
    dash[f"I{r}"].font = Font(name=FONT, bold=True, color=RED)
dash.cell(FL + N_FLAG_ROWS + 1, 2,
          f'=IF({FLAG_COUNT}>{N_FLAG_ROWS},"+ "&({FLAG_COUNT}-{N_FLAG_ROWS})&" more, see Colleagues sheet (Status = FLAG)",'
          f'IF({FLAG_COUNT}=0,"No one flagged yet.",""))')
dash.cell(FL + N_FLAG_ROWS + 1, 2).font = Font(name=FONT, italic=True, color=GREY)
protect(dash)

# ---------------------------------------------------------------- Macro Setup
title(setup, "MACRO SETUP (ONE TIME ONLY)",
      "Needs desktop Excel (Windows or Mac). Macros do not run in Excel Online or on phones.", 3)
setup.column_dimensions["A"].width = 6
setup.column_dimensions["B"].width = 95
setup_steps = [
    "Open this file in desktop Excel.",
    "Press Alt + F11 to open the Visual Basic editor (Mac: Tools > Macro > Visual Basic Editor).",
    "In the editor: File > Import File...  and choose GloveTracker.bas. Then close the editor.",
    "Press Alt + F8, choose SetupButtons, click Run. The red buttons appear on the Dashboard.",
    "File > Save As > choose 'Excel Macro-Enabled Workbook (*.xlsm)' and save.",
    "From now on open the .xlsm file. If Excel shows a yellow bar, click 'Enable Content'.",
]
header_row(setup, 4, ["Step", "What to do"])
for i, s in enumerate(setup_steps):
    r = 5 + i
    setup.cell(r, 1, i + 1)
    setup.cell(r, 2, s)
    body_cell(setup.cell(r, 1), center=True)
    body_cell(setup.cell(r, 2))
    setup.row_dimensions[r].height = 22
notes = [
    "Buttons: ISSUE GLOVES, UNDO LAST ISSUE, ADD COLLEAGUE (all on the Dashboard).",
    "Sheets are locked (no password) so entries go through the buttons. Pink/white boxes can still be typed in.",
    "To fix something by hand: Review > Unprotect Sheet. Remember to adjust Stock if you edit the Issue Log.",
]
for i, s in enumerate(notes):
    setup.cell(13 + i, 2, s).font = Font(name=FONT, italic=True, color=GREY)

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
dash.sheet_view.zoomScale = 90
wb.active = 0
wb.save(OUT)
print("saved", OUT)
